from g4f import Client
import docx

client  = Client()

def generate_code(prompt):
    response = client.chat.completions.create(
        model="gpt-4",
        messages=[{"role": "system", "content": "You are a helpful assistant that writes C programs."},
                  {"role": "user", "content": prompt}]
    )
    return response.choices[0].message.content

questions = [
    "Write a C program to display a linear representation of the sparse matrix.",
    "Write a C program for the matrix representation of polynomial equations.",
    "Write a C program to implement a stack with a maximum size of 10 with basic stack operations like push, pop and display.",
    "Write a C program to implement a stack with a maximum size of 10 with stack operations- peep and change.",
    "Write a C program using stack which will check that the given string belongs to grammar L= {wcwR | w {a, b}*} (Where wR is the reverse of w and c is in middle) or not.",
    "Write a C program using a stack to determine if an input character string is of the form ai bi where i >= 1. (You can use stack library)."
]

doc = docx.Document()
doc.add_heading('Lab Assignment Solutions', 0)

for i, question in enumerate(questions, start=1):
    doc.add_heading(f'Question {i}:', level=1)
    doc.add_paragraph(question)
    
    code_solution = generate_code(question)
    doc.add_heading('Solution:', level=2)
    doc.add_paragraph(code_solution, style='Normal')
    
    doc.add_heading('Test Cases:', level=2)
    test_cases = f"Provide at least four test cases for the following program: {question}"
    test_cases_solution = generate_code(test_cases)
    doc.add_paragraph(test_cases_solution, style='Normal')

file_path = "Lab_Assignment_3_Solutions.docx"
doc.save(file_path)
print(f"Document saved as {file_path}")